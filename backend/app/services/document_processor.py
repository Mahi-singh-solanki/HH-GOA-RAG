from __future__ import annotations

import csv
import json
import os
import shutil
from pathlib import Path
from typing import Any

import fitz
import pytesseract
from PIL import Image
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.config import get_settings


settings = get_settings()


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".docx",
    ".pptx",
    ".xlsx",
    ".csv",
    ".txt",
    ".md",
}


def ensure_directories() -> None:

    Path(settings.documents_dir).mkdir(
        parents=True,
        exist_ok=True,
    )

    Path(settings.qdrant_local_path).mkdir(
        parents=True,
        exist_ok=True,
    )

    Path(settings.pageindex_storage_path).mkdir(
        parents=True,
        exist_ok=True,
    )


def configure_tesseract() -> None:

    configured = settings.tesseract_cmd

    if configured:
        if Path(configured).exists():
            pytesseract.pytesseract.tesseract_cmd = configured
            return

    detected = shutil.which("tesseract")

    if detected:
        pytesseract.pytesseract.tesseract_cmd = detected
        return

    windows_default = Path(
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )

    if windows_default.exists():
        pytesseract.pytesseract.tesseract_cmd = str(
            windows_default
        )


configure_tesseract()


def clean_text(text: str) -> str:

    lines = []

    for line in text.splitlines():

        line = " ".join(line.split())

        if line:
            lines.append(line)

    return "\n".join(lines).strip()


def ocr_image(image: Image.Image) -> str:

    text = pytesseract.image_to_string(
        image,
        lang=settings.ocr_language,
    )

    return clean_text(text)


def extract_pdf(
    file_path: Path,
) -> tuple[list[dict[str, Any]], bool]:

    pages = []
    ocr_used = False

    document = fitz.open(file_path)

    try:

        for index, page in enumerate(document):

            page_number = index + 1

            native_text = clean_text(
                page.get_text("text")
            )

            if native_text:

                pages.append(
                    {
                        "page_number": page_number,
                        "text": native_text,
                        "source": "native_text",
                        "ocr_used": False,
                        "section": None,
                    }
                )

                continue

            if not settings.ocr_enabled:

                pages.append(
                    {
                        "page_number": page_number,
                        "text": "",
                        "source": "empty",
                        "ocr_used": False,
                        "section": None,
                    }
                )

                continue

            pixmap = page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False,
            )

            image = Image.frombytes(
                "RGB",
                [
                    pixmap.width,
                    pixmap.height,
                ],
                pixmap.samples,
            )

            text = ocr_image(image)

            ocr_used = True

            pages.append(
                {
                    "page_number": page_number,
                    "text": text,
                    "source": "ocr",
                    "ocr_used": True,
                    "section": None,
                }
            )

    finally:
        document.close()

    return pages, ocr_used


def extract_image(
    file_path: Path,
) -> tuple[list[dict[str, Any]], bool]:

    image = Image.open(file_path)

    text = ocr_image(image)

    return [
        {
            "page_number": 1,
            "text": text,
            "source": "ocr",
            "ocr_used": True,
            "section": None,
        }
    ], True


def extract_docx(
    file_path: Path,
) -> tuple[list[dict[str, Any]], bool]:

    document = Document(file_path)

    parts = []

    current_section = None

    for paragraph in document.paragraphs:

        text = clean_text(paragraph.text)

        if not text:
            continue

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        if "Heading" in style_name:

            current_section = text

            parts.append(
                f"[SECTION] {text}"
            )

        else:

            if current_section:
                parts.append(
                    f"[{current_section}] {text}"
                )
            else:
                parts.append(text)

    for table_index, table in enumerate(
        document.tables,
        start=1,
    ):

        rows = []

        for row in table.rows:

            cells = [
                clean_text(cell.text)
                for cell in row.cells
            ]

            rows.append(
                " | ".join(cells)
            )

        table_text = "\n".join(rows)

        parts.append(
            f"[TABLE {table_index}]\n"
            f"{table_text}"
        )

    text = "\n".join(parts)

    return [
        {
            "page_number": 1,
            "text": clean_text(text),
            "source": "docx",
            "ocr_used": False,
            "section": current_section,
        }
    ], False


def extract_pptx(
    file_path: Path,
) -> tuple[list[dict[str, Any]], bool]:

    presentation = Presentation(file_path)

    pages = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        parts = []

        title = None

        for shape in slide.shapes:

            if not hasattr(shape, "text"):
                continue

            text = clean_text(shape.text)

            if not text:
                continue

            if (
                getattr(shape, "is_placeholder", False)
                and shape.placeholder_format.type == 1
            ):
                title = text

            parts.append(text)

        if title:
            parts.insert(
                0,
                f"[SECTION] {title}",
            )

        pages.append(
            {
                "page_number": slide_number,
                "text": "\n".join(parts),
                "source": "pptx",
                "ocr_used": False,
                "section": title,
            }
        )

    return pages, False


def extract_xlsx(
    file_path: Path,
) -> tuple[list[dict[str, Any]], bool]:

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    pages = []

    for sheet_number, sheet in enumerate(
        workbook.worksheets,
        start=1,
    ):

        rows = []

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is None:
                    values.append("")
                else:
                    values.append(str(value))

            if any(values):
                rows.append(
                    " | ".join(values)
                )

        text = "\n".join(rows)

        pages.append(
            {
                "page_number": sheet_number,
                "text": text,
                "source": "xlsx",
                "ocr_used": False,
                "section": sheet.title,
            }
        )

    workbook.close()

    return pages, False


def extract_csv(
    file_path: Path,
) -> tuple[list[dict[str, Any]], bool]:

    with file_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.reader(file)

        rows = []

        for row in reader:

            rows.append(
                " | ".join(
                    str(value)
                    for value in row
                )
            )

    return [
        {
            "page_number": 1,
            "text": "\n".join(rows),
            "source": "csv",
            "ocr_used": False,
            "section": None,
        }
    ], False


def extract_text_file(
    file_path: Path,
) -> tuple[list[dict[str, Any]], bool]:

    text = file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    )

    return [
        {
            "page_number": 1,
            "text": clean_text(text),
            "source": file_path.suffix.lower().lstrip("."),
            "ocr_used": False,
            "section": None,
        }
    ], False


def process_document(
    file_path: Path,
    original_filename: str | None = None,
) -> dict[str, Any]:

    ensure_directories()

    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file extension: {extension}"
        )

    original_filename = (
        original_filename
        or file_path.name
    )

    if extension == ".pdf":

        pages, ocr_used = extract_pdf(
            file_path
        )

    elif extension in {
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
    }:

        pages, ocr_used = extract_image(
            file_path
        )

    elif extension == ".docx":

        pages, ocr_used = extract_docx(
            file_path
        )

    elif extension == ".pptx":

        pages, ocr_used = extract_pptx(
            file_path
        )

    elif extension == ".xlsx":

        pages, ocr_used = extract_xlsx(
            file_path
        )

    elif extension == ".csv":

        pages, ocr_used = extract_csv(
            file_path
        )

    else:

        pages, ocr_used = extract_text_file(
            file_path
        )

    document = {
        "document_id": file_path.stem,
        "filename": original_filename,
        "stored_filename": file_path.name,
        "file_type": extension.lstrip("."),
        "processing": {
            "ocr_enabled": settings.ocr_enabled,
            "ocr_used": ocr_used,
        },
        "pages": pages,
    }

    json_path = file_path.with_suffix(
        ".json"
    )

    json_path.write_text(
        json.dumps(
            document,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    return document